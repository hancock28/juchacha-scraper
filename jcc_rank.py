#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
剧查查小程序 · 三个漫剧榜单采集
===============================

实录协议（playlet-applet.dataeye.com，GET，免登录）：
  红果漫剧榜    /playlet/getComicsPlayletByDate?day=&pageId=&pageSize=&type=1   # 1 热播榜 / 2 新剧榜
  3D动态漫榜    /playlet/motionComic?day=&pageId=&pageSize=&rankType=1          # 1 动态漫 / 2 真人AI / 3 沙雕漫
  AI真人榜      /playlet/motionComic?day=&pageId=&pageSize=&rankType=2
日期来源：/playlet/getComicsPlayletDate?type=1   /playlet/motionComicDate

rankType 语义是用 /playlet/rankHonor 逐条回查反推出来的（day=2026-10-25 实测）：
  rankType=1 → 榜单荣誉 type 5（动态漫榜）  rankNum 1/2/3 与列表位置完全一致
  rankType=2 → type 7（真人AI榜）
  rankType=3 → type 6（沙雕漫榜）

用法：
  python3 jcc_rank.py                # 用接口返回的最新日期
  python3 jcc_rank.py --day 2026-10-25
  python3 jcc_rank.py --pages 5      # 每榜最多抓几页
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sqlite3
import sys
import time
import warnings
from typing import Any, Dict, List

warnings.filterwarnings("ignore")
import requests  # noqa: E402

BASE = "https://playlet-applet.dataeye.com"
UA = "Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) MicroMessenger/8.0.50"

RANKS = {
    # 名称(输出用)          path                             固定参数
    "红果漫剧榜": ("/playlet/getComicsPlayletByDate", {"type": 1}),
    "3D动态漫榜": ("/playlet/motionComic", {"rankType": 1}),
    "AI真人榜": ("/playlet/motionComic", {"rankType": 2}),
}
DATE_API = {
    "红果漫剧榜": ("/playlet/getComicsPlayletDate", {"type": 1}),
    "3D动态漫榜": ("/playlet/motionComicDate", {}),
    "AI真人榜": ("/playlet/motionComicDate", {}),
}


def get(path: str, params: Dict[str, Any], retry: int = 3) -> Dict[str, Any]:
    for i in range(retry):
        try:
            r = requests.get(BASE + path, params=params, headers={"User-Agent": UA},
                             timeout=25)
            js = r.json()
            if js.get("statusCode") == 200:
                return js
            raise RuntimeError(f"{path} -> {json.dumps(js, ensure_ascii=False)[:200]}")
        except Exception:
            if i == retry - 1:
                raise
            time.sleep(1.5 * (i + 1))
    return {}


def latest_day(name: str) -> str:
    path, fixed = DATE_API[name]
    js = get(path, dict(fixed))
    c = js.get("content") or {}
    return c.get("day") or ""


def fetch_rank(name: str, day: str, pages: int, size: int) -> List[Dict[str, Any]]:
    """实测：pageSize 硬上限 30；红果漫剧榜 pageId 1~3 有效（服务端第 4 页起回落重复），
    motionComic（动态漫/真人AI）分页无效、固定返回 Top30。因此做重复检测 + 去重。"""
    path, fixed = RANKS[name]
    size = min(size, 30)
    rows: List[Dict[str, Any]] = []
    seen = set()
    total = None
    for page in range(1, max(2, pages) + 1):
        js = get(path, {**fixed, "day": day, "pageId": page, "pageSize": size})
        pg = js.get("page") or {}
        if pg.get("totalRecords") is not None:
            total = pg["totalRecords"]
        chunk = js.get("content") or []
        if not chunk:
            break
        new = 0
        for r in chunk:
            pid = r.get("playletId")
            if pid in seen:
                continue
            seen.add(pid)
            r["_rank_name"] = name
            r["_rank_day"] = r.get("rankingDate") or r.get("calculateDate") or day
            r["_page"] = page
            rows.append(r)
            new += 1
        print(f"  {name} page{page}: +{new} 新 (累计 {len(rows)}"
              + (f"/{total}" if total else "") + ")")
        if new == 0:                      # 整页重复 = 分页到头
            break
        if total and len(rows) >= total:
            break
        time.sleep(0.6)
    return rows


def flatten(row: Dict[str, Any]) -> Dict[str, Any]:
    out = {}
    for k, v in row.items():
        if isinstance(v, list):
            if v and isinstance(v[0], dict):
                out[k] = "、".join(str(x.get("name") or x.get("publisherName") or "") for x in v)
            else:
                out[k] = "、".join(map(str, v))
        elif isinstance(v, str) and v.startswith("[") and v.endswith("]"):
            try:
                arr = json.loads(v)
                out[k] = "、".join(map(str, arr)) if isinstance(arr, list) else v
            except Exception:
                out[k] = v
        else:
            out[k] = v
    return out


def save(rows: List[Dict[str, Any]], outdir: str, day: str):
    os.makedirs(outdir, exist_ok=True)
    # 原始 JSON
    rawp = os.path.join(outdir, f"jcc_ranks_{day}.raw.json")
    with open(rawp, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=1)
    # 分榜 CSV
    by_rank: Dict[str, List[Dict[str, Any]]] = {}
    for r in rows:
        by_rank.setdefault(r["_rank_name"], []).append(flatten(r))
    for name, rs in by_rank.items():
        p = os.path.join(outdir, f"{name}_{day}.csv")
        cols: List[str] = []
        for r in rs:
            for k in r:
                if k not in cols:
                    cols.append(k)
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
            w.writeheader()
            for r in rs:
                w.writerow({k: r.get(k) for k in cols})
    # 合并 CSV
    allp = os.path.join(outdir, f"jcc_ranks_all_{day}.csv")
    cols = []
    for r in rows:
        fr = flatten(r)
        for k in fr:
            if k not in cols:
                cols.append(k)
    with open(allp, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(flatten(r))
    # SQLite
    dbp = os.path.join(outdir, "jcc_ranks.db")
    con = sqlite3.connect(dbp)
    con.execute("""CREATE TABLE IF NOT EXISTS rank_rows (
        rank_name TEXT, rank_day TEXT, ranking INTEGER, playlet_id INTEGER,
        playlet_name TEXT, native_id TEXT, hot_value INTEGER, hot_value_text TEXT,
        play_count TEXT, play_count_add TEXT, total_episode INTEGER,
        tags TEXT, producers TEXT, contractors TEXT, platform TEXT,
        cover TEXT, list_days INTEGER, raw TEXT,
        PRIMARY KEY (rank_name, rank_day, playlet_id) ON CONFLICT REPLACE)""")
    n = 0
    for r in rows:
        f = flatten(r)
        con.execute("INSERT INTO rank_rows VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (
            r.get("_rank_name"), r.get("_rank_day"), r.get("ranking"), r.get("playletId"),
            f.get("playletName"), r.get("playletNativeId"), r.get("hotValue"),
            r.get("hotValueText"), f.get("playCount"), f.get("playCountAdd"),
            r.get("totalEpisode"), f.get("playletTags") or f.get("contentTypes"),
            f.get("producerCompany") or f.get("producerList"),
            f.get("contractorCompany") or f.get("contractorList"),
            f.get("platformList"), r.get("cover"), r.get("listDays"),
            json.dumps(r, ensure_ascii=False)))
        n += 1
    con.commit()
    con.close()
    return rawp, allp, by_rank, dbp, n


def export_xlsx(by_rank: Dict[str, List[Dict[str, Any]]], outdir: str, day: str) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return ""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rs in by_rank.items():
        ws = wb.create_sheet(name[:28])
        cols: List[str] = []
        for r in rs:
            for k in r:
                if k not in cols:
                    cols.append(k)
        if not cols:
            continue
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rs:
            ws.append([r.get(k) for k in cols])
        ws.freeze_panes = "A2"
        for i, col in enumerate(cols, 1):
            width = min(40, max(10, len(str(col)) * 2))
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    p = os.path.join(outdir, f"剧查查三榜_{day}.xlsx")
    wb.save(p)
    return p


def summarize(by_rank: Dict[str, List[Dict[str, Any]]], top: int = 10):
    for name, rs in by_rank.items():
        rs = sorted(rs, key=lambda x: (x.get("ranking") or 10 ** 6))
        print(f"\n【{name}】共 {len(rs)} 部，Top{min(top, len(rs))}：")
        for r in rs[:top]:
            hot = r.get("hotValueText") or r.get("playCountAdd") or r.get("playCount") or "-"
            tags = r.get("playletTags") or r.get("contentTypes") or ""
            prod = r.get("producerCompany") or r.get("producerList") or ""
            print(f"  {str(r.get('ranking')):>3}. {str(r.get('playletName'))[:26]:28s} "
                  f"热度={hot:<8s} 集数={str(r.get('totalEpisode') or '-'):>5s}  "
                  f"标签={str(tags)[:24]:26s} 出品={str(prod)[:20]}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--day", default="", help="榜单日期 YYYY-MM-DD，默认取接口最新")
    ap.add_argument("--pages", type=int, default=6, help="最多翻几页（服务端 pageSize 上限 30）")
    ap.add_argument("--size", type=int, default=100)
    ap.add_argument("--out", default="data/jcc")
    a = ap.parse_args()

    days = {n: a.day or latest_day(n) for n in RANKS}
    for n, d in days.items():
        print(f"{n}: day={d or '(空)'}")
    rows: List[Dict[str, Any]] = []
    for name in RANKS:
        d = days[name]
        if not d:
            print(f"跳过 {name}：无可用日期")
            continue
        print(f"\n>> 抓取 {name} ({d})")
        rows.extend(fetch_rank(name, d, a.pages, a.size))

    if not rows:
        sys.exit("没抓到数据")

    day = max(days.values())
    rawp, allp, by_rank, dbp, n = save(rows, a.out, day)
    xlsx = export_xlsx(by_rank, a.out, day)
    summarize(by_rank)
    print(f"\n共 {n} 行")
    print("原始 JSON:", rawp)
    print("合并 CSV :", allp)
    for name in by_rank:
        print(f"分榜 CSV : {os.path.join(a.out, f'{name}_{day}.csv')}")
    if xlsx:
        print("Excel    :", xlsx)
    print("SQLite   :", dbp, "(table: rank_rows)")


if __name__ == "__main__":
    main()
